"""Tailored CVs: ask once about skills a job names, remember yes/no, and add confirmed skills to a per-job CV copy.

Only confirmed skills (status yes) are ever added; 'no' is stored so it is never asked again or added.
The CV's Word original (profile/cv-source/<CV file name>.docx, or the .docx next to the CV PDF) is never edited:
build() copies it, adds the skills to the matching line of its Skills section, exports base and copy to PDF with the
same converter (Microsoft Word, LibreOffice or docx2pdf), checks the page count did not change, and writes
data/inbox/cvs/<job_id>/<CV file name>. Anything that fails a check writes nothing.
"""

import importlib.util
import platform
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from . import formfill
from .db import ROOT
from .prescore import ALIASES

SOURCE_DIR = ROOT / "profile" / "cv-source"
WORK = ROOT / "data" / "cv-build"
# Word can only write inside its own sandbox folder; with Word the build works there and copies the PDF out.
WORD_WORK = Path.home() / "Library/Containers/com.microsoft.Word/Data/cv-build"
WORD_APPS = (Path("/Applications/Microsoft Word.app"), Path.home() / "Applications/Microsoft Word.app")
SOFFICE_PATHS = (
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "C:/Program Files/LibreOffice/program/soffice.exe", "C:/Program Files (x86)/LibreOffice/program/soffice.exe",
    "/usr/bin/soffice", "/usr/bin/libreoffice", "/usr/local/bin/soffice", "/usr/lib/libreoffice/program/soffice",
    "/opt/libreoffice/program/soffice", "/snap/bin/libreoffice",
)
INSTALL_HINT = ("Install one of: Microsoft Word (macOS), LibreOffice (any OS, https://www.libreoffice.org), "
                "or on Windows with Word: `uv sync --extra windows` (docx2pdf).")
MAX_ADD = 6  # anti-stuffing: never add more than this many skills to one CV
STATUSES = ("yes", "no")
LEVELS = ("basic", "working", "strong")
NO_SKILL = "no confirmed skill with a bullet on this CV"

# The Skills section is the lines after a "Skills" heading (in English or a few other languages) up to the next
# heading; each line is one skills group ("Languages: Go, Rust", "Python (Django, Flask)", or a single skill).
SKILLS_HEADING = re.compile(r"^(((technical|key|core|hard|it|professional|digital) )?skills|competenc(i|ies|e)s?|"
                            r"kompetenzen|fähigkeiten|kenntnisse|umiejętności|compétences|habilidades|competencias|"
                            r"competenze|vaardigheden|vaardigheid|competências)\b", re.I)
SECTION_HEADING = re.compile(r"^(education|(work |professional )?experience|employment|projects?|certifications?|"
                             r"languages|interests|hobbies|achievements|awards|courses|publications|volunteering|"
                             r"summary|profile|about( me)?|references|contact)\b", re.I)
# Skill families: a confirmed skill goes into the Skills line that already names the most of its relatives.
FAMILIES = (
    ("programming", "languages", "python", "java", "javascript", "typescript", "kotlin", "scala", "go", "golang",
     "rust", "swift", "c++", "c#", "php", "ruby", "dart", "perl"),
    ("backend", "django", "flask", "fastapi", "spring", "spring boot", ".net", "node.js", "express", "rest",
     "graphql", "microservices"),
    ("frontend", "react", "angular", "vue", "next.js", "html", "css", "tailwind", "redux"),
    ("cloud", "devops", "docker", "kubernetes", "linux", "ci/cd", "aws", "azure", "gcp", "terraform", "ansible",
     "jenkins", "deployment", "github actions", "gitlab ci"),
    ("data", "sql", "databases", "database", "postgresql", "mysql", "sqlite", "mongodb", "redis", "spark", "pandas",
     "numpy", "hadoop", "databricks", "power bi", "tableau", "excel", "etl", "analytics", "kafka"),
    ("testing", "qa", "selenium", "cypress", "playwright", "pytest", "junit", "jest", "postman"),
    ("marketing", "seo", "sem", "google ads", "meta ads", "google analytics", "looker studio", "hubspot", "salesforce",
     "crm", "mailchimp", "braze", "social media", "tiktok", "instagram", "copywriting", "wordpress", "shopify"),
    ("design", "canva", "photoshop", "illustrator", "indesign", "figma", "premiere pro", "after effects"),
    ("office", "excel", "powerpoint", "google sheets", "jira", "sap"),
)

# Terms prescore.ALIASES does not know yet (it has no Java, C#, etc.), and the tools of non-IT roles.
EXTRA_TERMS = {
    "java": [r"\bjava\b"], "spring boot": [r"spring boot", r"\bspring\b"], "kotlin": [r"kotlin"],
    "c#": [r"\bc#", r"\.net\b", r"asp\.net"], "c++": [r"\bc\+\+"], "scala": [r"\bscala\b"],
    "php": [r"\bphp\b"], "ruby": [r"\bruby\b"], "azure": [r"\bazure\b"], "aws": [r"\baws\b", r"amazon web services"],
    "gcp": [r"\bgcp\b", r"google cloud"], "kubernetes": [r"kubernetes", r"\bk8s\b"], "terraform": [r"terraform"],
    "power bi": [r"power ?bi"], "tableau": [r"tableau"], "excel": [r"\bexcel\b"], "pandas": [r"pandas"],
    "hadoop": [r"hadoop"], "spark": [r"\bspark\b", r"pyspark"], "databricks": [r"databricks"],
    "seo": [r"\bseo\b"], "sem": [r"\bsem\b"], "google ads": [r"google ads", r"adwords"],
    "meta ads": [r"meta ads", r"facebook ads", r"meta business"], "google analytics": [r"google analytics", r"\bga4\b"],
    "looker studio": [r"looker studio", r"data studio"], "hubspot": [r"hubspot"], "salesforce": [r"salesforce"],
    "crm": [r"\bcrm\b"], "mailchimp": [r"mailchimp"], "braze": [r"\bbraze\b"], "tiktok": [r"tiktok"],
    "instagram": [r"instagram"], "copywriting": [r"copywriting"], "wordpress": [r"wordpress"],
    "shopify": [r"shopify"], "canva": [r"\bcanva\b"], "photoshop": [r"photoshop"],
    "illustrator": [r"adobe illustrator", r"\billustrator\b"], "indesign": [r"indesign"], "figma": [r"\bfigma\b"],
    "premiere pro": [r"premiere pro"], "after effects": [r"after effects"], "powerpoint": [r"powerpoint"],
    "google sheets": [r"google sheets"], "jira": [r"\bjira\b"], "sap": [r"\bsap\b"],
}
TERMS = {**EXTRA_TERMS, **ALIASES}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def norm(skill: str) -> str:
    return re.sub(r"\s+", " ", skill.strip().lower())


def save_fact(conn, skill: str, status: str, display: str | None = None,
              level: str | None = None, section: str | None = None) -> dict:
    """Upsert one yes/no fact. Re-saving the same skill overwrites it."""
    if status not in STATUSES:
        raise ValueError(f"status must be one of {STATUSES}")
    if status == "yes" and level not in (None, *LEVELS):
        raise ValueError(f"level must be one of {LEVELS}")
    key = norm(skill)
    conn.execute(
        "INSERT INTO skill_facts (skill, display, status, level, section_key, updated_at) VALUES (?,?,?,?,?,?) "
        "ON CONFLICT(skill) DO UPDATE SET display=excluded.display, status=excluded.status, "
        "level=excluded.level, section_key=excluded.section_key, updated_at=excluded.updated_at",
        (key, display or skill.strip(), status, level if status == "yes" else None, section, _now()),
    )
    conn.commit()
    return {"skill": key, "status": status, "display": display or skill.strip(), "level": level, "section": section}


def list_facts(conn) -> list[dict]:
    rows = conn.execute("SELECT skill, display, status, level, section_key, updated_at FROM skill_facts "
                        "ORDER BY skill").fetchall()
    return [dict(r) for r in rows]


def _facts(conn) -> dict[str, dict]:
    return {f["skill"]: f for f in list_facts(conn)}


# ---------- the CV's Word original and its Skills section ----------

def docx_candidates(pdf: Path, cv: str = "main", source_dir: Path | None = None) -> list[Path]:
    """Where the Word original of a CV may be, in lookup order: <source_dir>/<variant>/<name>.docx (variants only),
    <source_dir>/<name>.docx, or the .docx next to the PDF. <name> is the PDF's file name."""
    source_dir = source_dir or SOURCE_DIR
    return ([source_dir / cv / f"{pdf.stem}.docx"] if cv != "main" else []) + \
        [source_dir / f"{pdf.stem}.docx", pdf.with_suffix(".docx")]


def source_docx(cv: str = "main") -> Path:
    """The Word original of a CV variant: the first candidate that exists, else where it is expected."""
    candidates = docx_candidates(formfill.cv_path(cv), cv)
    return next((c for c in candidates if c.exists()), candidates[-2])


def _open(cv: str):
    """(source path, python-docx Document or None when the original is missing)."""
    src = source_docx(cv)
    if not src.exists():
        return src, None
    import docx  # python-docx (project dependency)
    return src, docx.Document(src)


def _is_heading(paragraph) -> bool:
    text = paragraph.text.strip()
    if not text or len(text) > 40 or "," in text:
        return False
    style = getattr(paragraph.style, "name", "") or ""
    return (style.startswith(("Heading", "Title")) or text.isupper()
            or bool(SECTION_HEADING.match(text) or SKILLS_HEADING.match(text)))


def _key(text: str) -> str:
    """'Cloud & DevOps: Docker, Linux' -> 'cloud-devops'; 'Python (Django, Flask)' -> 'python'."""
    label = text.split(":", 1)[0] if ":" in text[:40] else re.split(r"\s*[(,]", text, maxsplit=1)[0]
    return re.sub(r"[^a-z0-9+#]+", "-", label.lower()).strip("-")[:30] or "skills"


def layout(doc) -> dict[str, dict]:
    """Skills lines of the CV: {key: {index, mode, text}}. mode "inside_parens" adds inside the last "( ... )" when
    the line ends with ")", "append" adds ", X" at the end."""
    sections: dict[str, dict] = {}
    inside = False
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if _is_heading(p):
            inside = bool(SKILLS_HEADING.match(text))
            continue
        if not inside or not text:
            continue
        key = base = _key(text)
        n = 2
        while key in sections:
            key, n = f"{base}-{n}", n + 1
        sections[key] = {"index": i, "mode": "inside_parens" if text.endswith(")") else "append", "text": text}
    return sections


def _mentions(text: str, word: str) -> bool:
    return re.search(rf"(?<![\w+#]){re.escape(word)}(?![\w+#])", text.lower()) is not None


def guess_section(skill: str, sections: dict[str, dict]) -> str | None:
    """The Skills line naming the most relatives of the skill (FAMILIES); None when no line names any."""
    related = {w for family in FAMILIES if skill in family for w in family} - {skill}
    best, best_hits = None, 0
    for key, spec in sections.items():
        hits = sum(1 for w in related if _mentions(spec["text"], w))
        if hits > best_hits:
            best, best_hits = key, hits
    return best


def cv_skills(cv: str = "main", doc=None) -> set[str]:
    """Skills already on the CV: the profile's skill list plus the tech terms the Word original names."""
    profile = formfill.load_profile()
    names = {norm(s["name"]) for s in profile.get("skills") or [] if isinstance(s, dict) and s.get("name")}
    if doc is not None:
        names |= job_terms({"description": "\n".join(p.text for p in doc.paragraphs)})
    return names


def job_terms(job: dict) -> set[str]:
    """Canonical terms the job names: its skills list plus tech terms found in title and description."""
    found = {norm(s) for s in (job.get("skills") or []) if s}
    text = f"{job.get('title') or ''}\n{(job.get('description') or '')[:6000]}".lower()
    for term, patterns in TERMS.items():
        if any(re.search(pat, text) for pat in patterns):
            found.add(term)
    return found


def plan(conn, job: dict, cv: str = "main") -> dict:
    """Split the job's terms into: already on CV, confirmed to add, denied, and never asked (unknown).
    Confirmed skills with no Skills line to go into (or no Word original) land in no_bullet."""
    facts = _facts(conn)
    src, doc = _open(cv)
    sections = layout(doc) if doc is not None else {}
    on_cv = cv_skills(cv, doc)
    buckets = {"on_cv": [], "known_add": [], "denied": [], "unknown": [], "no_bullet": []}
    for term in sorted(job_terms(job)):
        if term in on_cv:
            buckets["on_cv"].append(term)
        elif term in facts and facts[term]["status"] == "yes":
            section = facts[term]["section_key"]
            if section not in sections:
                section = guess_section(term, sections)
            item = {"skill": term, "display": facts[term]["display"], "level": facts[term]["level"], "section": section}
            (buckets["known_add"] if section else buckets["no_bullet"]).append(item)
        elif term in facts:
            buckets["denied"].append(term)
        else:
            buckets["unknown"].append(term)
    if len(buckets["known_add"]) > MAX_ADD:
        buckets["deferred_over_cap"] = [k["skill"] for k in buckets["known_add"][MAX_ADD:]]
        buckets["known_add"] = buckets["known_add"][:MAX_ADD]
    return {"job_id": job.get("id"), "cv": cv, "source": str(src), "source_found": doc is not None,
            "sections": {k: v["text"] for k, v in sections.items()}, **buckets}


def tailored_path(job_id: int, cv: str = "main") -> Path:
    return formfill.CVS / str(job_id) / formfill.cv_path(cv).name


# ---------- DOCX editing ----------

def _add_to_bullet(paragraph, items: list[str], mode: str) -> None:
    """Edit only the line's last text run (a bold label run before it stays as it is)."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        raise ValueError(f"line {paragraph.text[:30]!r} has no text run; refusing to edit")
    run = runs[-1]
    text = run.text.rstrip()
    extra = ", ".join(items)
    if mode == "append":
        run.text = f"{text}, {extra}"
    elif mode == "inside_parens":
        if not text.endswith(")"):
            raise ValueError(f"line {paragraph.text[:30]!r} does not end with ')'")
        run.text = f"{text[:-1]}, {extra})"
    else:
        raise ValueError(f"unknown mode {mode}")


def apply_skills(doc, additions: list[dict], sections: dict[str, dict] | None = None) -> dict[str, list[str]]:
    """Insert each addition into its Skills line. Returns {section: [display names added]}."""
    sections = layout(doc) if sections is None else sections
    grouped: dict[str, list[str]] = {}
    for item in additions:
        grouped.setdefault(item["section"], []).append(item["display"])
    for section, items in grouped.items():
        if section not in sections:
            raise ValueError(f"no Skills line {section!r} on this CV; have {sorted(sections)}")
        spec = sections[section]
        _add_to_bullet(doc.paragraphs[spec["index"]], items, spec["mode"])
    return grouped


def _check_only_additions(before: list[str], after: list[str], added: list[str]) -> None:
    """Every changed paragraph must differ only by the added skill names (no other text touched)."""
    if len(before) != len(after):
        raise ValueError("paragraph count changed")
    expected = set(added)
    for old, new in zip(before, after):
        if old == new:
            continue
        old_tokens = set(re.findall(r"[^,()]+", old))
        new_tokens = set(re.findall(r"[^,()]+", new))
        extra = {t.strip() for t in new_tokens - old_tokens}
        missing = {t.strip() for t in old_tokens - new_tokens}
        if missing or not extra <= {e.strip() for e in expected} | {""}:
            raise ValueError(f"unexpected change: {old!r} -> {new!r}")


# ---------- PDF export: Word (macOS) -> LibreOffice (any OS) -> docx2pdf (Windows) ----------

def _word_available() -> bool:
    return platform.system() == "Darwin" and any(app.exists() for app in WORD_APPS)


def _word_export_pdf(docx: Path, pdf: Path) -> None:
    """Export with Microsoft Word, the program most CV PDFs come from. Raises on failure."""
    name = docx.name
    # try/on error: a failed save must still close the document, or every later export finds it open and fails.
    lines = [
        'tell application "Microsoft Word"',
        f'open POSIX file "{docx}"',
        "delay 2",
        "try",
        f'save as document "{name}" file name "{pdf}" file format format PDF',
        "on error msg",
        f'close document "{name}" saving no',
        "error msg",
        "end try",
        f'close document "{name}" saving no',
        "end tell",
    ]
    cmd = ["osascript"] + [arg for line in lines for arg in ("-e", line)]
    res = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if res.returncode != 0 or not pdf.exists() or pdf.stat().st_size == 0:
        raise RuntimeError(f"Word export failed: {res.stderr.strip() or 'no PDF written'}")


def find_soffice() -> str | None:
    """LibreOffice's command line program: on PATH or in its usual install folder."""
    for name in ("soffice", "libreoffice"):
        if path := shutil.which(name):
            return path
    for path in SOFFICE_PATHS:
        if Path(path).exists():
            return path
    return next((str(p) for p in sorted(Path("/opt").glob("libreoffice*/program/soffice"))), None)


def _soffice_export_pdf(docx: Path, pdf: Path) -> None:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("LibreOffice not found")
    with tempfile.TemporaryDirectory() as tmp:
        # A private user profile: a running LibreOffice window would otherwise swallow the headless call.
        cmd = [soffice, f"-env:UserInstallation={(Path(tmp) / 'profile').as_uri()}", "--headless",
               "--convert-to", "pdf", "--outdir", tmp, str(docx)]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        out = Path(tmp) / f"{docx.stem}.pdf"
        if res.returncode != 0 or not out.exists() or out.stat().st_size == 0:
            raise RuntimeError(f"LibreOffice export failed: {res.stderr.strip() or 'no PDF written'}")
        shutil.move(str(out), str(pdf))


def _docx2pdf_available() -> bool:
    return platform.system() == "Windows" and importlib.util.find_spec("docx2pdf") is not None


def _docx2pdf_export_pdf(docx: Path, pdf: Path) -> None:
    from docx2pdf import convert  # optional extra: uv sync --extra windows
    convert(str(docx), str(pdf))
    if not pdf.exists() or pdf.stat().st_size == 0:
        raise RuntimeError("docx2pdf wrote no PDF")


def exporters() -> list[tuple[str, object]]:
    """Available PDF exporters, most faithful first."""
    chain = [("word", _word_available, _word_export_pdf),
             ("libreoffice", lambda: find_soffice() is not None, _soffice_export_pdf),
             ("docx2pdf", _docx2pdf_available, _docx2pdf_export_pdf)]
    return [(name, run) for name, available, run in chain if available()]


def pdf_exporter() -> str | None:
    """Name of the exporter build() would use first, None when there is none."""
    return next((name for name, _ in exporters()), None)


def test_export() -> str:
    """Export a one-line document the way build() does: proves the exporter really works (macOS may first ask to let
    the terminal control Word). Returns the exporter's name; raises RuntimeError with the reason."""
    import docx  # python-docx (project dependency)
    root = WORD_WORK if _word_available() else WORK
    root.mkdir(parents=True, exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="export-test-", dir=root))
    try:
        src = work / "export-test.docx"
        page = docx.Document()
        page.add_paragraph("Job Copilot export test")
        page.save(src)
        return export_pdf(src, work / "export-test.pdf")
    finally:
        shutil.rmtree(work, ignore_errors=True)


def export_pdf(docx: Path, pdf: Path, only: str | None = None) -> str:
    """Export with the first exporter that works (or only `only`, so base and copy match). Returns its name."""
    errors = []
    for name, run in exporters():
        if only and name != only:
            continue
        try:
            run(docx, pdf)
            return name
        except (RuntimeError, OSError, subprocess.SubprocessError) as e:
            errors.append(f"{name}: {e}")
    raise RuntimeError("could not export the CV to PDF (" + ("; ".join(errors) or "no exporter found") + "). "
                       + INSTALL_HINT)


def pdf_pages(pdf: Path) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes()))


# ---------- build ----------

def build(conn, job: dict, cv: str = "main") -> dict:
    """Produce the per-job CV copy. Writes data/inbox/cvs/<job_id>/ only if every check passes."""
    result = plan(conn, job, cv)
    additions = result["known_add"]
    if not additions and not result["no_bullet"]:
        return {**result, "built": False, "reason": NO_SKILL}
    if not result["source_found"]:
        return {**result, "built": False,
                "blocked": f"original CV source missing: put the Word (.docx) file of this CV at {result['source']}"}
    if not additions:
        return {**result, "built": False, "reason": NO_SKILL}

    root = WORD_WORK if _word_available() else WORK
    root.mkdir(parents=True, exist_ok=True)
    # A fresh folder per build, removed afterwards: CV copies must not pile up in Word's container (one folder per
    # job ever tailored did), and two projects building the same job id must not share one.
    work = Path(tempfile.mkdtemp(prefix=f"{job['id']}-{cv}-", dir=root))
    try:
        return _build_in(work, result, additions, job, cv)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def _build_in(work: Path, result: dict, additions: list[dict], job: dict, cv: str) -> dict:
    import docx  # python-docx (project dependency)

    base_docx, tailored_docx = work / "base.docx", work / "tailored.docx"
    base_pdf, tailored_pdf = work / "base.pdf", work / "tailored.pdf"
    shutil.copyfile(result["source"], base_docx)  # Word and python-docx work on copies only

    doc = docx.Document(base_docx)
    before = [p.text for p in doc.paragraphs]
    placed = apply_skills(doc, additions)
    after = [p.text for p in doc.paragraphs]
    added = [item["display"] for item in additions]
    _check_only_additions(before, after, added)
    doc.save(tailored_docx)

    try:
        exporter = export_pdf(base_docx, base_pdf)
        export_pdf(tailored_docx, tailored_pdf, only=exporter)
    except RuntimeError as e:
        return {**result, "built": False, "blocked": str(e)}
    if pdf_pages(tailored_pdf) != pdf_pages(base_pdf):
        return {**result, "built": False, "blocked": "page count changed; nothing written"}

    dest = tailored_path(job["id"], cv)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(tailored_pdf, dest)
    return {**result, "built": True, "added": added, "placed": placed, "pages": pdf_pages(dest),
            "exporter": exporter, "pdf": str(dest)}
