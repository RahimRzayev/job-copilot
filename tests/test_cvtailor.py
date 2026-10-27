from pathlib import Path

import pytest

from jobhunt import cvtailor, formfill
from jobhunt.db import connect

FIXTURE_PROFILE = Path(__file__).parent / "fixtures" / "profile.yaml"
SKILLS_CV = ["Alex Example", "SKILLS", "Python (Flask, Pandas)", "Languages: Kotlin, C, Ruby",
             "Cloud & DevOps: Docker, Ansible, CI/CD, Monitoring", "SQL & Databases (MySQL, MongoDB)",
             "EXPERIENCE", "Set up laptops, answered tickets, wrote docs"]


def make_docx(path: Path, lines: list[str]) -> Path:
    import docx
    doc = docx.Document()
    for line in lines:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


@pytest.fixture(autouse=True)
def profile(tmp_path, monkeypatch):
    """Fixture profile, Word originals in tmp_path (cv.docx with a Skills section, cv-data.docx without one)."""
    monkeypatch.setattr(formfill, "PROFILE", FIXTURE_PROFILE)
    monkeypatch.setattr(formfill, "CVS", tmp_path / "cvs")
    monkeypatch.setattr(cvtailor, "SOURCE_DIR", tmp_path / "cv-source")
    make_docx(tmp_path / "cv-source" / "cv.docx", SKILLS_CV)
    make_docx(tmp_path / "cv-source" / "cv-data.docx", ["Alex Example", "Data analyst, reporting, dashboards"])


@pytest.fixture
def conn(tmp_path):
    return connect(tmp_path / "test.db")


def job(text="", skills=None, title="Backend intern", id=1):
    return {"id": id, "title": title, "description": text, "skills": skills or []}


def test_save_fact_upserts_and_normalises(conn):
    cvtailor.save_fact(conn, "  Java ", "no")
    cvtailor.save_fact(conn, "java", "yes", display="Java", level="basic", section="backend")
    facts = cvtailor.list_facts(conn)
    assert len(facts) == 1
    assert facts[0]["skill"] == "java" and facts[0]["status"] == "yes"
    assert facts[0]["display"] == "Java" and facts[0]["section_key"] == "backend"


def test_save_fact_rejects_bad_status_and_level(conn):
    with pytest.raises(ValueError):
        cvtailor.save_fact(conn, "java", "maybe")
    with pytest.raises(ValueError):
        cvtailor.save_fact(conn, "java", "yes", level="expert")


def test_layout_reads_the_skills_section_only(tmp_path):
    import docx
    sections = cvtailor.layout(docx.Document(tmp_path / "cv-source" / "cv.docx"))
    assert {k: v["mode"] for k, v in sections.items()} == {
        "python": "inside_parens", "languages": "append", "cloud-devops": "append", "sql-databases": "inside_parens"}


def test_guess_section_uses_skill_families(tmp_path):
    import docx
    sections = cvtailor.layout(docx.Document(tmp_path / "cv-source" / "cv.docx"))
    assert cvtailor.guess_section("fastapi", sections) == "python"
    assert cvtailor.guess_section("java", sections) == "languages"
    assert cvtailor.guess_section("terraform", sections) == "cloud-devops"
    assert cvtailor.guess_section("tableau", sections) == "sql-databases"
    assert cvtailor.guess_section("jira", sections) is None


def test_yes_fact_is_added_and_no_fact_is_denied(conn):
    cvtailor.save_fact(conn, "java", "yes", display="Java", level="basic", section="languages")
    cvtailor.save_fact(conn, "scala", "no")
    result = cvtailor.plan(conn, job("We use Java and Scala for services."))
    assert [k["skill"] for k in result["known_add"]] == ["java"]
    assert result["denied"] == ["scala"]
    assert "scala" not in result["unknown"]
    assert result["source_found"] and "languages" in result["sections"]


def test_unknown_skill_is_reported_for_asking(conn):
    result = cvtailor.plan(conn, job("Experience with Java required."))
    assert "java" in result["unknown"]
    assert result["known_add"] == []


def test_skill_already_on_cv_is_not_asked_or_added(conn):
    # Python is in the profile's skills and Ruby is on the Word original: neither is unknown or an addition.
    result = cvtailor.plan(conn, job("Python, Flask and Ruby experience."))
    assert {"python", "ruby"} <= set(result["on_cv"])
    assert "python" not in result["unknown"] and all(k["skill"] != "python" for k in result["known_add"])


def test_additions_are_capped(conn):
    names = ["java", "c++", "c#", "aws", "azure", "gcp", "kubernetes", "spark"]  # each has a line on the CV
    for n in names:
        cvtailor.save_fact(conn, n, "yes", display=n.title(), level="basic")
    result = cvtailor.plan(conn, job(" ".join(names)))
    assert len(result["known_add"]) == cvtailor.MAX_ADD
    assert len(result["deferred_over_cap"]) == len(names) - cvtailor.MAX_ADD


def test_build_writes_nothing_without_confirmed_skill(conn):
    out = cvtailor.build(conn, job("Java services."))
    assert out["built"] is False and out["reason"] == "no confirmed skill with a bullet on this CV"


def test_build_refuses_when_source_cv_missing(conn, monkeypatch, tmp_path):
    cvtailor.save_fact(conn, "java", "yes", display="Java", level="basic", section="languages")
    monkeypatch.setattr(cvtailor, "SOURCE_DIR", tmp_path / "nowhere")
    out = cvtailor.build(conn, job("Java services."))
    assert out["built"] is False and "source" in out["blocked"]
    assert not cvtailor.tailored_path(1).exists()


def test_build_writes_tailored_pdf_with_the_cv_file_name(conn, monkeypatch, tmp_path):
    def fake_export(docx_path, pdf, only=None):
        pdf.write_bytes(b"%PDF-1.4 /Type /Page\n")
        return "libreoffice"
    monkeypatch.setattr(cvtailor, "_word_available", lambda: False)
    monkeypatch.setattr(cvtailor, "WORK", tmp_path / "work")
    monkeypatch.setattr(cvtailor, "export_pdf", fake_export)
    cvtailor.save_fact(conn, "aws", "yes", display="AWS", level="basic")
    out = cvtailor.build(conn, job("AWS services.", id=7))
    assert out["built"] and out["placed"] == {"cloud-devops": ["AWS"]}
    assert Path(out["pdf"]) == tmp_path / "cvs" / "7" / "cv.pdf" and Path(out["pdf"]).exists()
    assert list((tmp_path / "work").iterdir()) == []  # no CV copies left behind


def test_formfill_uses_base_cv_when_no_tailored_copy():
    data = formfill.build_data({"id": 987654321, "title": "x", "company": "Acme", "source": "linkedin", "skills": []})
    assert data["resume_path"].endswith("cv.pdf")
    assert "cvs" not in Path(data["resume_path"]).parts


def test_apply_skills_inserts_into_mapped_bullets_only(tmp_path):
    import docx

    doc = docx.Document(tmp_path / "cv-source" / "cv.docx")
    before = [p.text for p in doc.paragraphs]
    additions = [{"skill": "java", "display": "Java", "section": "languages"},
                 {"skill": "aws", "display": "AWS", "section": "cloud-devops"},
                 {"skill": "spark", "display": "Spark", "section": "sql-databases"}]
    placed = cvtailor.apply_skills(doc, additions)
    after = [p.text for p in doc.paragraphs]
    cvtailor._check_only_additions(before, after, ["Java", "AWS", "Spark"])
    assert placed == {"languages": ["Java"], "cloud-devops": ["AWS"], "sql-databases": ["Spark"]}
    assert "Languages: Kotlin, C, Ruby, Java" in after
    assert any(t.endswith("Monitoring, AWS") for t in after)
    assert any("MongoDB, Spark)" in t for t in after)


def test_apply_skills_keeps_a_bold_label_run():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Skills")
    line = doc.add_paragraph()
    line.add_run("Tools: ").bold = True
    line.add_run("Git, Jira")
    cvtailor.apply_skills(doc, [{"skill": "docker", "display": "Docker", "section": "tools"}])
    assert [r.text for r in doc.paragraphs[1].runs] == ["Tools: ", "Git, Jira, Docker"]


def test_check_rejects_unexpected_change():
    before = ["Kotlin, C, Ruby", "Summary text"]
    after = ["Kotlin, C, Ruby, Java", "Summary text changed"]
    with pytest.raises(ValueError):
        cvtailor._check_only_additions(before, after, ["Java"])


def test_plan_sends_unplaceable_skill_to_no_bullet(conn):
    # The data-analyst CV has no Skills section, so confirmed skills are reported, not added.
    cvtailor.save_fact(conn, "java", "yes", display="Java", level="basic")
    result = cvtailor.plan(conn, job("Java services."), cv="data-analyst")
    assert result["known_add"] == [] and [k["skill"] for k in result["no_bullet"]] == ["java"]


def test_pdf_export_falls_back_to_libreoffice_without_word(monkeypatch, tmp_path):
    calls = []

    def fake_run(cmd, **kwargs):
        calls.append(cmd)
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        (outdir / "cv.pdf").write_bytes(b"%PDF-1.4 /Type /Page\n")
        return type("Done", (), {"returncode": 0, "stderr": ""})()

    monkeypatch.setattr(cvtailor, "_word_available", lambda: False)
    monkeypatch.setattr(cvtailor, "_docx2pdf_available", lambda: False)
    monkeypatch.setattr(cvtailor, "find_soffice", lambda: "/opt/fake/soffice")
    monkeypatch.setattr(cvtailor.subprocess, "run", fake_run)
    pdf = tmp_path / "out.pdf"
    assert cvtailor.pdf_exporter() == "libreoffice"
    assert cvtailor.export_pdf(tmp_path / "cv-source" / "cv.docx", pdf) == "libreoffice"
    assert pdf.exists() and calls[0][0] == "/opt/fake/soffice" and "--headless" in calls[0]


def test_pdf_export_without_any_exporter_names_the_options(monkeypatch, tmp_path):
    monkeypatch.setattr(cvtailor, "_word_available", lambda: False)
    monkeypatch.setattr(cvtailor, "_docx2pdf_available", lambda: False)
    monkeypatch.setattr(cvtailor, "find_soffice", lambda: None)
    assert cvtailor.pdf_exporter() is None
    with pytest.raises(RuntimeError, match="LibreOffice"):
        cvtailor.export_pdf(tmp_path / "cv-source" / "cv.docx", tmp_path / "out.pdf")


def test_non_it_tools_are_recognised_and_excellent_is_not_excel():
    terms = cvtailor.job_terms({"title": "Marketing Assistant", "description": (
        "Excellent writing. HubSpot, Google Ads, Meta Ads, SEO, Photoshop, WordPress, TikTok and Instagram. "
        "Customer reactivation campaigns in Braze.")})
    assert {"hubspot", "google ads", "meta ads", "seo", "photoshop", "wordpress", "tiktok", "instagram",
            "braze"} <= terms
    assert not terms & {"excel", "react"}


def test_marketing_skills_go_to_the_marketing_line(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Skills", 1)
    d.add_paragraph("Marketing: social media, copywriting, Mailchimp, Google Analytics")
    d.add_paragraph("Design: Canva")
    sections = cvtailor.layout(d)
    assert cvtailor.guess_section("hubspot", sections) == "marketing"
    assert cvtailor.guess_section("photoshop", sections) == "design"
